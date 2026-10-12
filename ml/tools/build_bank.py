# -*- coding: utf-8 -*-
"""Бодлогын санг SymPy-аар шалгаад бодлогын_сан.xlsx, CSV болгон гаргана.
Ажиллуулах: python3 build_bank.py   (sympy, openpyxl шаардлагатай)"""
import csv
import sys
from pathlib import Path

from sympy import *  # noqa: F401,F403
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bank_data import GLOSSARY, P, TOPICS
from bank_steps import STEPS

x, y, a, b, k = symbols("x y a b k")
HERE = Path(__file__).resolve().parent
OUT = HERE.parent  # capstone_datasets/


def _num(t):
    return float(N(t))


def same(g, e):
    if isinstance(g, dict):
        return {str(kk): simplify(v) for kk, v in g.items()} == {str(kk): simplify(v) for kk, v in e.items()}
    if isinstance(g, (list, tuple)):
        if len(g) != len(e):
            return False
        return all(simplify(u - v) == 0 for u, v in zip(sorted(g, key=_num), sorted(e, key=_num)))
    if isinstance(g, Set):
        return g == e
    return simplify(g - e) == 0


def main():
    rows, fails = [], []
    for i, p in enumerate(P, 1):
        (topic, sub, lvl, text, latex, ans, check, exp, terms, hints, mis) = p
        g, e = eval(check), eval(exp)
        ok = same(g, e)
        if not ok:
            fails.append((i, latex, g, e))
        pid = f"{topic}-{i:03d}"
        prereq = dict((t[0], t[3]) for t in TOPICS)[topic]
        steps = " || ".join(st for st in STEPS.get(pid, []) if st)  # FR-07 бодолтын алхмууд
        rows.append([pid, topic, sub, lvl, text, latex, ans, str(e), prereq, terms,
                     hints[0], hints[1], hints[2], mis, "OK" if ok else "ШАЛГАХ", steps])
    print(f"SymPy шалгалт: {len(P) - len(fails)}/{len(P)} зөв")
    for f in fails:
        print("  ЗӨРҮҮ:", f)

    head = ["id", "сэдэв_код", "дэд_сэдэв", "түвшин(1-3)", "бодлогын_текст", "томьёо_LaTeX",
            "хариу", "хариу_sympy", "prerequisite", "нэр_томьёо", "сануулга_1", "сануулга_2",
            "сануулга_3", "түгээмэл_алдаа", "sympy_шалгалт", "бодолтын_алхам"]

    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")

    def sheet(ws, header, data, widths):
        ws.append(header)
        for r in data:
            ws.append(r)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Бодлого"
    sheet(ws, head, rows, [11, 9, 22, 9, 40, 30, 18, 18, 12, 30, 38, 38, 38, 38, 10, 60])

    topic_name = {t[0]: t[1] for t in TOPICS}
    ws2 = wb.create_sheet("Сэдвүүд")
    trows = [[c, n, g, pr, " ; ".join(topic_name[q] for q in pr.split(";") if q),
              sum(1 for p in P if p[0] == c)] for c, n, g, pr in TOPICS]
    sheet(ws2, ["код", "сэдэв", "анги", "prerequisite_код", "prerequisite_нэр", "бодлогын_тоо"], trows,
          [8, 30, 7, 16, 45, 13])

    ws3 = wb.create_sheet("Сэдвийн_граф")
    edges = [[pr, c, topic_name[pr], topic_name[c]] for c, n, g, prs in TOPICS for pr in prs.split(";") if pr]
    sheet(ws3, ["эх (өмнөх)", "төгсгөл (дараах)", "эх_нэр", "төгсгөл_нэр"], edges, [12, 15, 30, 30])

    ws4 = wb.create_sheet("Нэр_томьёо")
    sheet(ws4, ["нэр_томьёо", "энгийн_тайлбар", "жишээ", "сэдэв_код"], [list(g) for g in GLOSSARY],
          [28, 60, 34, 10])

    xlsx = OUT / "бодлогын_сан.xlsx"
    wb.save(xlsx)
    csv_dir = OUT / "csv"
    csv_dir.mkdir(exist_ok=True)
    for name, h, d in [("problems", head, rows),
                       ("topics", ["code", "name", "grade", "prereq"], [list(t) for t in TOPICS]),
                       ("topic_edges", ["from", "to", "from_name", "to_name"], edges),
                       ("glossary", ["term", "definition", "example", "topic"], [list(g) for g in GLOSSARY])]:
        with open(csv_dir / f"{name}.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(h)
            w.writerows(d)
    print("Бичсэн:", xlsx.name, "+ csv/")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
